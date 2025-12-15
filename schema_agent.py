"""
智能Schema生成系统 - 文档分析与Schema生成
"""

import os
import json
from typing import Dict, List, Any, Optional, Tuple
from datetime import datetime
from langchain_core.messages import SystemMessage, HumanMessage
import re

# 导入LLM服务模块
from llm_service import LLMService, get_llm_service, create_llm_service


class DocumentAnalysisReport:
    """文档分析报告类"""
    
    def __init__(self):
        self.domain = "未识别"
        self.sub_domain = ""
        self.keywords = []
        self.entities = []
        self.relations = []
        self.attributes = []
        self.inferred_entities = []
        self.inferred_relations = []
        self.confidence = 0.0
        self.document_summary = ""
        self.raw_content = ""
        
    def to_dict(self) -> Dict:
        """转换为字典格式"""
        return {
            "domain": self.domain,
            "sub_domain": self.sub_domain,
            "keywords": self.keywords,
            "entities": self.entities,
            "relations": self.relations,
            "attributes": self.attributes,
            "inferred_entities": self.inferred_entities,
            "inferred_relations": self.inferred_relations,
            "confidence": self.confidence,
            "document_summary": self.document_summary
        }
    
    def from_dict(self, data: Dict):
        """从字典加载"""
        self.domain = data.get("domain", "未识别")
        self.sub_domain = data.get("sub_domain", "")
        self.keywords = data.get("keywords", [])
        self.entities = data.get("entities", [])
        self.relations = data.get("relations", [])
        self.attributes = data.get("attributes", [])
        self.inferred_entities = data.get("inferred_entities", [])
        self.inferred_relations = data.get("inferred_relations", [])
        self.confidence = data.get("confidence", 0.0)
        self.document_summary = data.get("document_summary", "")
        return self
    
    def to_markdown(self) -> str:
        """生成Markdown格式报告"""
        md = f"""# 文档分析报告

## 1. 领域识别

- **主要领域**: {self.domain}
- **子领域**: {self.sub_domain}
- **置信度**: {self.confidence:.2%}

## 2. 领域关键词

{', '.join(self.keywords) if self.keywords else '未识别到关键词'}

## 3. 文档摘要

{self.document_summary}

## 4. 识别的实体类型

| 实体类型 | 中文名称 | 描述 | 置信度 |
|---------|---------|------|--------|
"""
        for entity in self.entities:
            md += f"| {entity.get('name', '')} | {entity.get('cn_name', '')} | {entity.get('description', '')} | {entity.get('confidence', 0):.2%} |\n"
        
        md += """
## 5. 识别的关系类型

| 关系类型 | 中文名称 | 源实体 | 目标实体 | 置信度 |
|---------|---------|--------|---------|--------|
"""
        for relation in self.relations:
            md += f"| {relation.get('name', '')} | {relation.get('cn_name', '')} | {relation.get('source', '')} | {relation.get('target', '')} | {relation.get('confidence', 0):.2%} |\n"
        
        md += """
## 6. 识别的属性

| 属性名 | 中文名称 | 数据类型 | 所属实体 | 示例值 |
|--------|---------|---------|---------|--------|
"""
        for attr in self.attributes:
            md += f"| {attr.get('name', '')} | {attr.get('cn_name', '')} | {attr.get('data_type', '')} | {attr.get('entity', '')} | {attr.get('example', '')} |\n"
        
        md += """
## 7. 推测的实体（需要确认）

| 实体类型 | 推测理由 | 建议操作 |
|---------|---------|---------|
"""
        for entity in self.inferred_entities:
            md += f"| {entity.get('name', '')} | {entity.get('reason', '')} | {entity.get('suggestion', '')} |\n"
        
        md += """
## 8. 推测的关系（需要确认）

| 关系类型 | 推测理由 | 建议操作 |
|---------|---------|---------|
"""
        for relation in self.inferred_relations:
            md += f"| {relation.get('name', '')} | {relation.get('reason', '')} | {relation.get('suggestion', '')} |\n"
        
        return md


class SchemaAgent:
    """智能Schema生成Agent"""
    
    def __init__(self, model_config_id: str = None):
        """
        初始化Agent
        
        Args:
            model_config_id: 模型配置ID，如果为None则使用默认配置
        """
        # 使用LLM服务
        self.llm_service = create_llm_service(model_config_id)
        self.current_report = None
    
    def set_model(self, model_config_id: str = None) -> bool:
        """
        设置使用的模型
        
        Args:
            model_config_id: 模型配置ID，如果为None则使用默认配置
            
        Returns:
            是否设置成功
        """
        return self.llm_service.set_model(model_config_id)
    
    def get_current_model_info(self) -> str:
        """获取当前模型信息"""
        return self.llm_service.get_model_display_name()
    
    def get_available_models(self) -> List[tuple]:
        """获取所有可用的模型选项"""
        return self.llm_service.get_available_models()
    
    def get_default_model_id(self) -> Optional[str]:
        """获取默认模型ID"""
        return self.llm_service.get_default_model_id()
    
    @property
    def llm(self):
        """获取底层LLM实例（兼容旧代码）"""
        return self.llm_service.llm
        
    def _get_file_extension(self, file_path: str) -> str:
        """获取文件扩展名（小写）"""
        return os.path.splitext(file_path)[1].lower()
    
    def _convert_to_markdown(self, file_path: str) -> Tuple[str, str]:
        """
        将文档转换为Markdown格式
        
        Args:
            file_path: 文档路径
            
        Returns:
            (Markdown内容, Markdown文件路径)
        """
        file_ext = self._get_file_extension(file_path)
        file_name_base = os.path.splitext(os.path.basename(file_path))[0]
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
        # 对于需要转换的文档类型
        if file_ext in ['.doc', '.docx', '.pdf']:
            print(f"📄 检测到 {file_ext} 文件，需要转换为Markdown格式...")
            
            # 读取二进制文件内容
            with open(file_path, 'rb') as f:
                binary_content = f.read()
            
            # 根据文件类型生成不同的prompt
            if file_ext == '.pdf':
                file_type_desc = "PDF文档"
                additional_instructions = """
- 保留所有的标题层级结构
- 表格转换为Markdown表格格式
- 列表项保持原有的层级关系
- 忽略页眉页脚和页码
- 保留所有重要的文本内容
"""
            elif file_ext in ['.doc', '.docx']:
                file_type_desc = "Word文档"
                additional_instructions = """
- 保留所有的标题层级结构（使用#号表示层级）
- 表格转换为Markdown表格格式
- 有序列表和无序列表保持原格式
- 保留文档中的重要格式（如加粗、斜体）
- 忽略页眉页脚、水印等非正文内容
"""
            else:
                file_type_desc = "文档"
                additional_instructions = ""
            
            # 使用LLM转换文档
            prompt = f"""
请将以下{file_type_desc}的内容转换为清晰的Markdown格式。

转换要求：
{additional_instructions}
- 确保内容结构清晰、层次分明
- 使用恰当的Markdown语法
- 保留所有重要信息，不要遗漏关键内容

注意：由于我无法直接读取二进制文件，请基于文档的一般结构和内容特征，
生成一个结构化的Markdown文档。如果是合同、报告等正式文档，
请包含标题、章节、条款、表格等典型元素。

文件名：{os.path.basename(file_path)}
文件大小：{len(binary_content)} 字节

请生成相应的Markdown格式内容：
"""
            
            # 对于实际的doc/docx/pdf文件，尝试提取文本
            extracted_text = self._try_extract_text(file_path, file_ext)
            if extracted_text:
                prompt = f"""
请将以下{file_type_desc}的内容转换为清晰的Markdown格式。

原始文档内容：
{extracted_text[:5000]}  # 只使用前5000字符

转换要求：
{additional_instructions}
- 确保内容结构清晰、层次分明
- 使用恰当的Markdown语法
- 保留所有重要信息
- 如果内容被截断，请在最后注明

请生成完整的Markdown格式内容：
"""
            
            messages = [
                SystemMessage(content="你是一个专业的文档格式转换专家，擅长将各种格式的文档转换为Markdown格式"),
                HumanMessage(content=prompt)
            ]
            
            response = self.llm.invoke(messages)
            markdown_content = response.content
            
            # 保存转换后的Markdown文件
            markdown_file_path = os.path.join(
                "data/input",
                f"{timestamp}_{file_name_base}.md"
            )
            os.makedirs(os.path.dirname(markdown_file_path), exist_ok=True)
            
            with open(markdown_file_path, 'w', encoding='utf-8') as f:
                f.write(markdown_content)
            
            print(f"✅ 文档已转换为Markdown格式并保存到: {markdown_file_path}")
            
            return markdown_content, markdown_file_path
            
        # 对于直接支持的格式
        elif file_ext in ['.csv', '.txt', '.md']:
            print(f"✅ 检测到 {file_ext} 文件，直接进行分析...")
            content = self._read_document_with_encoding(file_path)
            
            # 对于CSV文件，可以转换为表格格式的Markdown
            if file_ext == '.csv':
                content = self._csv_to_markdown(content)
            
            return content, file_path
        
        else:
            # 未知格式，尝试作为文本读取
            print(f"⚠️ 未知文件格式 {file_ext}，尝试作为文本文件处理...")
            content = self._read_document_with_encoding(file_path)
            return content, file_path
    
    def _try_extract_text(self, file_path: str, file_ext: str) -> Optional[str]:
        """
        尝试从文档中提取文本
        
        Args:
            file_path: 文件路径
            file_ext: 文件扩展名
            
        Returns:
            提取的文本内容，如果失败返回None
        """
        try:
            if file_ext in ['.doc', '.docx']:
                # 尝试使用python-docx库
                try:
                    from docx import Document
                    doc = Document(file_path)
                    paragraphs = []
                    for para in doc.paragraphs:
                        if para.text.strip():
                            paragraphs.append(para.text)
                    
                    # 也提取表格
                    for table in doc.tables:
                        for row in table.rows:
                            row_text = []
                            for cell in row.cells:
                                row_text.append(cell.text.strip())
                            if any(row_text):
                                paragraphs.append(' | '.join(row_text))
                    
                    return '\n'.join(paragraphs)
                except ImportError:
                    print("⚠️ python-docx库未安装，无法直接读取Word文档")
                except Exception as e:
                    print(f"⚠️ 读取Word文档失败: {e}")
            
            elif file_ext == '.pdf':
                # 尝试使用PyPDF2或pdfplumber
                try:
                    import PyPDF2
                    with open(file_path, 'rb') as f:
                        pdf_reader = PyPDF2.PdfReader(f)
                        text_content = []
                        for page_num in range(len(pdf_reader.pages)):
                            page = pdf_reader.pages[page_num]
                            text_content.append(page.extract_text())
                        return '\n'.join(text_content)
                except ImportError:
                    try:
                        import pdfplumber
                        with pdfplumber.open(file_path) as pdf:
                            text_content = []
                            for page in pdf.pages:
                                text_content.append(page.extract_text())
                            return '\n'.join(text_content)
                    except ImportError:
                        print("⚠️ PyPDF2和pdfplumber库都未安装，无法直接读取PDF文档")
                except Exception as e:
                    print(f"⚠️ 读取PDF文档失败: {e}")
        
        except Exception as e:
            print(f"⚠️ 提取文档文本时出错: {e}")
        
        return None
    
    def _csv_to_markdown(self, csv_content: str) -> str:
        """
        将CSV内容转换为Markdown表格
        
        Args:
            csv_content: CSV文本内容
            
        Returns:
            Markdown格式的表格
        """
        lines = csv_content.strip().split('\n')
        if not lines:
            return csv_content
        
        # 简单的CSV解析（不处理引号等复杂情况）
        rows = [line.split(',') for line in lines]
        
        # 构建Markdown表格
        markdown_lines = []
        
        # 表头
        if rows:
            markdown_lines.append('| ' + ' | '.join(rows[0]) + ' |')
            markdown_lines.append('|' + '---|' * len(rows[0]))
            
            # 数据行
            for row in rows[1:]:
                # 确保每行有相同数量的列
                while len(row) < len(rows[0]):
                    row.append('')
                markdown_lines.append('| ' + ' | '.join(row[:len(rows[0])]) + ' |')
        
        return '\n'.join(markdown_lines)
    
    def _read_document_with_encoding(self, document_path: str) -> str:
        """尝试不同编码方式读取文档"""
        content = None
        encodings = ['utf-8', 'gbk', 'gb2312', 'gb18030', 'latin-1']
        
        for encoding in encodings:
            try:
                with open(document_path, 'r', encoding=encoding) as f:
                    content = f.read()
                print(f"✅ 使用 {encoding} 编码成功读取文件")
                break
            except UnicodeDecodeError:
                continue
            except Exception as e:
                print(f"读取文件时出错: {e}")
                continue
        
        if content is None:
            try:
                with open(document_path, 'rb') as f:
                    binary_content = f.read()
                content = binary_content.decode('utf-8', errors='ignore')
                print("⚠️ 使用忽略错误模式读取文件")
            except Exception as e:
                raise Exception(f"无法读取文件 {document_path}: {str(e)}")
        
        return content
    
    def stage1_analyze_document(self, document_path: str) -> DocumentAnalysisReport:
        """
        分析文档，生成分析报告
        
        Args:
            document_path: 文档路径
            
        Returns:
            文档分析报告
        """
        print("="*60)
        print("🔍 文档分析阶段")
        print("="*60)
        
        # 创建报告对象
        report = DocumentAnalysisReport()
        
        # 处理不同类型的文档
        content, processed_path = self._convert_to_markdown(document_path)
        report.raw_content = content[:3000]  # 保存前3000字作为参考
        
        # Step 1: 领域识别
        print("\n📌 步骤1：识别文档领域...")
        domain_info = self._identify_domain(content)
        report.domain = domain_info.get("domain", "未识别")
        report.sub_domain = domain_info.get("sub_domain", "")
        report.keywords = domain_info.get("keywords", [])
        report.confidence = domain_info.get("confidence", 0.0)
        
        # Step 2: 激活领域知识并分析文档
        print("\n📌 步骤2：基于领域知识分析文档结构...")
        analysis_result = self._analyze_with_domain_knowledge(content, domain_info)
        
        # 填充报告
        report.document_summary = analysis_result.get("summary", "")
        report.entities = analysis_result.get("entities", [])
        report.relations = analysis_result.get("relations", [])
        report.attributes = analysis_result.get("attributes", [])
        report.inferred_entities = analysis_result.get("inferred_entities", [])
        report.inferred_relations = analysis_result.get("inferred_relations", [])
        
        self.current_report = report
        
        print("\n✅ 文档分析完成!")
        print(f"  - 领域: {report.domain}")
        print(f"  - 实体类型: {len(report.entities)} 个")
        print(f"  - 关系类型: {len(report.relations)} 个")
        print(f"  - 属性: {len(report.attributes)} 个")
        
        return report
    
    def _identify_domain(self, content: str) -> Dict:
        """识别文档领域"""
        sample_content = content[:1500]
        
        prompt = f"""
你是一个专业的文档领域识别专家。请分析以下文档内容，识别它属于哪个专业领域。

文档内容（前1500字）：
{sample_content}

请识别并返回JSON格式：
{{
    "domain": "主要领域名称（如：供应链管理、项目管理、财务、法律、医疗、教育、制造业、电商等）",
    "sub_domain": "子领域名称",
    "keywords": ["关键词1", "关键词2", "关键词3", "..."],
    "features": ["特征1", "特征2"],
    "domain_description": "领域描述",
    "confidence": 0.95
}}
"""
        
        messages = [
            SystemMessage(content="你是文档领域识别专家"),
            HumanMessage(content=prompt)
        ]
        
        response = self.llm.invoke(messages)
        
        try:
            response_text = response.content
            json_match = re.search(r'\{[\s\S]*\}', response_text)
            if json_match:
                return json.loads(json_match.group())
        except Exception as e:
            print(f"解析领域识别结果失败: {e}")
        
        return {
            "domain": "通用",
            "sub_domain": "未知",
            "keywords": [],
            "confidence": 0.5
        }
    
    def _analyze_with_domain_knowledge(self, content: str, domain_info: Dict) -> Dict:
        """基于领域知识分析文档"""
        sample_content = content[:2500]
        
        domain = domain_info.get('domain', '通用')
        keywords = domain_info.get('keywords', [])
        
        prompt = f"""
你是{domain}领域的专家，正在分析该领域的文档。

领域背景：
- 主领域：{domain}
- 子领域：{domain_info.get('sub_domain', '')}
- 关键概念：{', '.join(keywords)}

文档内容（前2500字）：
{sample_content}

请完成以下任务：

1. 生成文档摘要（100-200字）
2. 识别文档中的实体类型（不是实体实例）
3. 识别实体之间的关系类型
4. 识别实体的属性
5. 推测可能存在但文档中未明确的实体和关系

输出JSON格式：
{{
    "summary": "文档摘要",
    "entities": [
        {{
            "name": "EntityName（英文PascalCase）",
            "cn_name": "实体中文名",
            "description": "实体描述",
            "confidence": 0.9,
            "evidence": ["文档中的证据"]
        }}
    ],
    "relations": [
        {{
            "name": "relationName（英文camelCase）",
            "cn_name": "关系中文名",
            "source": "源实体英文名",
            "target": "目标实体英文名",
            "description": "关系描述",
            "confidence": 0.85,
            "evidence": ["文档中的证据"]
        }}
    ],
    "attributes": [
        {{
            "name": "attributeName（英文camelCase）",
            "cn_name": "属性中文名",
            "entity": "所属实体英文名",
            "data_type": "Text/Float/Integer（仅支持这三种类型）",
            "example": "示例值",
            "description": "属性描述"
        }}
    ],
    "inferred_entities": [
        {{
            "name": "推测的实体名",
            "reason": "推测理由",
            "suggestion": "建议如何确认"
        }}
    ],
    "inferred_relations": [
        {{
            "name": "推测的关系名",
            "source": "源实体",
            "target": "目标实体",
            "reason": "推测理由",
            "suggestion": "建议如何确认"
        }}
    ]
}}

重要规则：
1. 实体名用英文PascalCase（如：Project, Supplier）
2. 关系和属性用英文camelCase（如：belongsTo, projectName）
3. 基于{domain}领域的专业知识进行分析
4. 区分确定的和推测的内容
"""
        
        messages = [
            SystemMessage(content=f"你是{domain}领域的资深专家，精通该领域的业务模型和数据结构"),
            HumanMessage(content=prompt)
        ]
        
        response = self.llm.invoke(messages)
        
        try:
            response_text = response.content
            json_match = re.search(r'\{[\s\S]*\}', response_text)
            if json_match:
                return json.loads(json_match.group())
        except Exception as e:
            print(f"解析分析结果失败: {e}")
        
        return {
            "summary": "",
            "entities": [],
            "relations": [],
            "attributes": [],
            "inferred_entities": [],
            "inferred_relations": []
        }
    
    def stage2_generate_schema(self, report: DocumentAnalysisReport, template_id: str = "openspg_kag_v1", 
                               previous_schema: str = None, validation_errors: List[Dict] = None) -> str:
        """
        基于分析报告生成Schema
        
        Args:
            report: 文档分析报告
            template_id: Schema模板ID
            previous_schema: 之前生成的Schema（用于重新生成）
            validation_errors: 验证错误信息（用于重新生成）
            
        Returns:
            生成的Schema文本
        """
        print("\n" + "="*60)
        if validation_errors:
            print("🔄 重新生成Schema（基于错误反馈）")
        else:
            print("🏗️ Schema生成阶段")
        print("="*60)
        
        # 构建增强的Schema生成prompt
        domain_context = f"""
## 领域背景
- 主领域：{report.domain}
- 子领域：{report.sub_domain}
- 关键概念：{', '.join(report.keywords)}
- 文档摘要：{report.document_summary}
"""
        
        # 准备实体、关系和属性信息 - 按实体组织
        entities_with_props = {}
        for entity in report.entities:
            entity_name = entity.get('name', '')
            entity_cn_name = entity.get('cn_name', '')
            entities_with_props[entity_name] = {
                'cn_name': entity_cn_name,
                'description': entity.get('description', ''),
                'properties': [],
                'relations': []
            }
            
            # 确保每个实体至少有name和description属性
            # 添加默认的name属性
            entities_with_props[entity_name]['properties'].append({
                'name': f'{entity_name.lower()}Name',
                'cn_name': f'{entity_cn_name}名称',
                'data_type': 'Text',
                'description': f'{entity_cn_name}的名称'
            })
            # 添加默认的description属性
            entities_with_props[entity_name]['properties'].append({
                'name': 'description',
                'cn_name': '描述',
                'data_type': 'Text',
                'description': f'{entity_cn_name}的详细描述'
            })
        
        # 分配属性到对应实体（避免重复）- 处理类型转换
        for attr in report.attributes:
            entity_name = attr.get('entity', '')
            attr_name = attr.get('name', '')
            data_type = attr.get('data_type', 'Text')
            
            # 转换不支持的类型
            if data_type not in ['Text', 'Float', 'Integer']:
                print(f"⚠️ 属性 {attr_name} 的类型 {data_type} 不支持，转换为Text")
                attr['data_type'] = 'Text'
                attr['original_type'] = data_type  # 记录原始类型
            
            if entity_name in entities_with_props:
                # 检查是否已经添加了同名属性（避免与默认属性冲突）
                existing_names = [p['name'] for p in entities_with_props[entity_name]['properties']]
                if attr_name not in existing_names and attr_name.lower() not in [n.lower() for n in existing_names]:
                    entities_with_props[entity_name]['properties'].append(attr)
        
        # 分配关系到源实体
        for relation in report.relations:
            source = relation.get('source', '')
            if source in entities_with_props:
                entities_with_props[source]['relations'].append(relation)
        
        entities_info = json.dumps(entities_with_props, indent=2, ensure_ascii=False)
        
        # 生成合适的namespace（全英文）
        namespace = report.domain.replace(' ', '').replace('管理', 'Management').replace('系统', 'System')
        namespace = namespace.replace('项目', 'Project').replace('供应链', 'SupplyChain')
        namespace = namespace.replace('医疗', 'Medical').replace('法律', 'Legal')
        namespace = namespace.replace('金融', 'Finance').replace('电商', 'Ecommerce')
        # 如果还有中文，使用默认namespace
        if any('\u4e00' <= c <= '\u9fa5' for c in namespace):
            namespace = 'AutoGenerated'
        
        # 如果有之前的Schema和错误信息，添加到prompt中
        error_context = ""
        if previous_schema and validation_errors:
            error_messages = []
            for error in validation_errors:
                if isinstance(error, dict):
                    msg = error.get('message', str(error))
                    details = error.get('details', '')
                    rule_name = error.get('rule_name', '')
                    if rule_name:
                        error_messages.append(f"[{rule_name}] {msg}")
                    else:
                        error_messages.append(msg)
                    if details:
                        error_messages.append(f"  详情: {details}")
                elif hasattr(error, 'message'):
                    error_messages.append(error.message)
                else:
                    error_messages.append(str(error))
            
            error_context = f"""
## ⚠️ 重要：之前生成的Schema存在以下验证错误，请务必修复：

### 错误列表：
{chr(10).join('- ' + msg for msg in error_messages)}

### 之前生成的Schema（有错误）：
```
{previous_schema}
```

请仔细分析这些错误，并在新生成的Schema中修复所有问题。确保：
1. 所有实体名称使用英文PascalCase格式（如Project, Supplier）
2. 所有属性和关系名使用英文camelCase格式（如projectName, belongsTo）
3. 不要使用下划线（_）
4. 每个实体都有properties部分
5. 只有当实体真正有关系时才添加relations部分
6. 属性类型只能是Text、Float或Integer（不支持Date、DateTime等其他类型）
7. 关系的目标实体必须在Schema中定义
"""
        
        prompt = f"""
你是一个Schema设计专家，擅长为{report.domain}-{report.sub_domain}领域设计数据模型。

{domain_context}

## 已确认的数据结构（按实体组织）：
{entities_info}

{error_context}

请基于以上信息，生成一个完整的Schema定义。

极其严格的格式要求（必须100%遵循）：
1. namespace必须是: {namespace}
2. 每个实体声明格式：EntityName(实体中文名): EntityType
3. 每个实体命名不能有特殊字符,"合同/主合同"是错误的，只能是"合同"或者"主合同"
4. properties: 必须缩进4个空格
5. 每个实体必须至少包含：
   - entityName(实体名称): Text
   - description(描述): Text
6. 每个属性格式：propertyNameInEnglish(属性中文名): Text、Float或Integer
   - 属性名必须是英文camelCase格式
   - 括号内是中文描述
   - 类型只能是Text、Float或Integer（日期类型应该用Text表示）
7. relations: 必须缩进4个空格（只有当实体有关系时才包含此部分）
8. 每个关系格式：relationNameInEnglish(关系中文名): TargetEntityName
   - 关系名必须是英文camelCase格式
   - 目标实体必须是英文名称

- 每个实体的完整定义必须是：
  ```
  EntityName(实体中文名): EntityType
      properties:
          propertyName(属性中文名): Type
          ...
      relations:  # 如果有关系才写这部分
          relationName(关系中文名): TargetEntity
          ...
  ```
  
⚠️ 极其重要的规则（违反将导致错误）：
- **每个实体只能定义一次**！同一个实体的properties和relations必须写在一起，如下正确示例。**
- **每个实体只能定义一次**！同一个实体的properties和relations必须写在一起，如下正确示例。**
- **每个实体只能定义一次**！同一个实体的properties和relations必须写在一起，如下正确示例。**


正确示例（必须这样写）：
```
namespace {namespace}

Contract(合同): EntityType
    properties:
        contractName(合同名称): Text
        description(描述): Text
        contractNumber(合同编号): Text
        totalAmount(总金额): Float
    relations:
        hasParty(有甲方): PartyA
        hasClause(包含条款): ContractClause

PartyA(甲方): EntityType
    properties:
        partyaName(甲方名称): Text
        description(描述): Text
        registrationNumber(注册号): Text

ContractClause(合同条款): EntityType
    properties:
        clauseName(条款名称): Text
        description(描述): Text
        clauseContent(条款内容): Text
```

错误示例：
```
Contract(合同): EntityType
    properties:
        contractName(合同名称): Text
        description(描述): Text

PartyA(甲方): EntityType
    properties:
        partyaName(甲方名称): Text

Contract(合同): EntityType  # ❌ 错误！Contract已经定义过了
    relations:
        hasParty(有甲方): PartyA
```



重要规则总结：
1. 每个实体必须完整定义在一个地方
2. 每个实体必须有properties部分，至少包含name和description
3. 只有当实体真正有关系时才包含relations部分
4. 如果实体没有关系，不要写relations:这一行
5. 每个实体只能定义一次
5. 每个实体只能定义一次

基于提供的数据结构信息生成Schema，严格按照格式要求。
"""
        
        messages = [
            SystemMessage(content=f"你是{report.domain}-{report.sub_domain}领域的Schema设计专家"),
            HumanMessage(content=prompt)
        ]
        
        response = self.llm.invoke(messages)
        schema_text = response.content
        
        # 清理输出（去除可能的markdown标记）
        schema_text = schema_text.replace('```yaml', '').replace('```', '').strip()
        
        print("\n✅ Schema生成完成!")
        
        return schema_text
    
    def edit_report(self, report_dict: Dict) -> DocumentAnalysisReport:
        """
        编辑分析报告
        
        Args:
            report_dict: 编辑后的报告字典
            
        Returns:
            更新后的报告对象
        """
        report = DocumentAnalysisReport()
        report.from_dict(report_dict)
        self.current_report = report
        return report
    
    def save_report(self, report: DocumentAnalysisReport, output_path: str = None, original_filename: str = None) -> str:
        """
        保存分析报告
        
        Args:
            report: 分析报告
            output_path: 输出路径
            original_filename: 原始文件名（用于生成标准化的输出文件名）
            
        Returns:
            保存的文件路径
        """
        if output_path is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            if original_filename:
                # 从原始文件名中提取基础名称（去掉扩展名）
                base_name = os.path.splitext(os.path.basename(original_filename))[0]
                output_path = f"data/output/report_{base_name}_{timestamp}.json"
            else:
                output_path = f"data/output/analysis_report_{timestamp}.json"
        
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(report.to_dict(), f, indent=2, ensure_ascii=False)
        
        # 同时保存Markdown版本
        md_path = output_path.replace('.json', '.md')
        with open(md_path, 'w', encoding='utf-8') as f:
            f.write(report.to_markdown())
        
        print(f"✅ 报告已保存到: {output_path}")
        print(f"✅ Markdown版本: {md_path}")
        
        return output_path
    
    def load_report(self, report_path: str) -> DocumentAnalysisReport:
        """
        加载分析报告
        
        Args:
            report_path: 报告文件路径
            
        Returns:
            分析报告对象
        """
        with open(report_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        report = DocumentAnalysisReport()
        report.from_dict(data)
        self.current_report = report
        
        return report
    
    def save_schema(self, schema: str, original_filename: str = None) -> str:
        """
        保存生成的Schema
        
        Args:
            schema: Schema文本
            original_filename: 原始文件名
            
        Returns:
            保存的文件路径
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
        if original_filename:
            base_name = os.path.splitext(os.path.basename(original_filename))[0]
            schema_path = f"data/output/schema_{base_name}_{timestamp}.yaml"
        else:
            schema_path = f"data/output/two_stage_schema_{timestamp}.yaml"
        
        os.makedirs(os.path.dirname(schema_path), exist_ok=True)
        
        with open(schema_path, 'w', encoding='utf-8') as f:
            f.write(schema)
        
        print(f"✅ Schema已保存到: {schema_path}")
        
        return schema_path


def test_schema_generation():
    """测试Schema生成"""
    print("\n" + "="*70)
    print("🚀 智能Schema生成测试")
    print("="*70)
    
    # 创建Agent
    agent = SchemaAgent()
    
    # 创建测试文档
    test_doc = "test_two_stage_doc.txt"
    with open(test_doc, 'w', encoding='utf-8') as f:
        f.write("""
项目管理系统数据说明

一、项目信息
项目编号：PROJ-2024-001
项目名称：智能供应链管理平台
项目经理：张三
项目状态：进行中
开始日期：2024-01-15
预算金额：5,000,000元
所属部门：信息技术部

二、供应商管理
供应商编号：SUP-001
供应商名称：北京科技有限公司
联系人：李四
联系电话：010-12345678
供应商类型：技术服务
注册资本：10,000,000元
合作等级：战略合作伙伴

三、采购订单
订单编号：PO-2024-0315
订单日期：2024-03-15
供应商：北京科技有限公司
采购金额：2,000,000元
交付日期：2024-04-15
订单状态：已确认

四、合同信息
合同编号：CONTRACT-2024-001
合同名称：技术服务合同
甲方：我方公司
乙方：北京科技有限公司
签订日期：2024-03-01
合同金额：3,000,000元
有效期：2024-03-01至2025-03-01

五、业务关系
1. 项目包含多个采购订单
2. 供应商签署合同
3. 采购订单关联供应商
4. 项目由部门管理
5. 合同约束供应商服务
        """)
    
    # 文档分析
    print("\n" + "-"*60)
    report = agent.stage1_analyze_document(test_doc)
    
    # 显示分析报告
    print("\n📊 分析报告预览：")
    print("-"*60)
    print(report.to_markdown())
    
    # 保存报告
    report_path = agent.save_report(report, original_filename=test_doc)
    
    # 模拟用户编辑（这里直接使用原报告）
    print("\n" + "-"*60)
    print("💡 用户可以在此编辑报告...")
    print("   (演示中跳过编辑步骤)")
    
    # 生成Schema
    print("\n" + "-"*60)
    schema = agent.stage2_generate_schema(report)
    
    print("\n📋 生成的Schema：")
    print("-"*60)
    print(schema)
    
    # 保存Schema
    schema_path = agent.save_schema(schema, original_filename=test_doc)
    
    return report, schema


if __name__ == "__main__":
    test_schema_generation()
